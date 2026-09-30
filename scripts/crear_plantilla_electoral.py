#!/usr/bin/env python3
"""Extrae Seguridad del machote electoral recibido y coloca sus variables."""

from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from editorial import configurar
from fuentes_word import incrustar
from migrar_machotes_v2 import save


ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT / 'config/variables_plataforma_electoral_seguridad.json'
LABELS = (
    ('Propuesta o plan general:', 'propuesta_plan_general'),
    ('Problema social a resolver:', 'problema_social'),
    ('Razonamiento justificatorio:', 'razonamiento_justificatorio'),
    ('Propuestas específicas:', 'propuestas_especificas'),
    ('Línea discursiva:', 'linea_discursiva'),
)


def list_level(paragraph):
    props = paragraph._p.pPr
    number = props.numPr if props is not None else None
    return number.ilvl.val if number is not None and number.ilvl is not None else None


def heading(paragraph, value):
    """Conserva las propiedades visuales de la fuente sin su numeración previa."""
    run_properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
        paragraph._p.pPr.remove(paragraph._p.pPr.numPr)
    paragraph.style = 'Normal'
    run = paragraph.add_run(value)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def replace_whole(paragraph, expected, value):
    if paragraph.text.strip() != expected:
        raise ValueError(f'Cambió el marcador de la fuente: {expected}')
    props = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(value)
    if props is not None:
        run._r.insert(0, props)


def replace_field(paragraph, label, marker):
    if paragraph.text.strip() != label + ' XXXX':
        raise ValueError(f'Cambió el campo del machote electoral: {label}')
    changed = 0
    for run in paragraph.runs:
        if 'XXXX' in run.text:
            run.text = run.text.replace('XXXX', marker)
            changed += 1
    if changed != 1:
        raise ValueError(f'El campo {label} no tiene un valor único.')


def main():
    spec = json.loads(SPEC.read_text(encoding='utf-8'))
    source = ROOT / spec['fuente_machote']
    if hashlib.sha256(source.read_bytes()).hexdigest() != spec['fuente_machote_sha256']:
        raise ValueError('Cambió el machote electoral recibido: revisar las correspondencias.')
    document = Document(source)
    paragraphs = document.paragraphs
    starts = [i for i, p in enumerate(paragraphs)
              if p.text.strip() == 'Seguridad' and list_level(p) == 0]
    if len(starts) != 1:
        raise ValueError('Se requiere un único capítulo Seguridad en el machote electoral.')
    start = starts[0]
    end = next((i for i in range(start + 1, len(paragraphs))
                if list_level(paragraphs[i]) == 0), len(paragraphs))
    chapter = paragraphs[start:end]
    children = [p for p in chapter if list_level(p) == 1]
    if [p.text.strip() for p in children] != ['Protección Civil', 'Seguridad']:
        raise ValueError('El capítulo debe contener Protección Civil y Seguridad, en ese orden.')

    replace_whole(paragraphs[0], 'PLATAFORMA ELECTORAL', 'Plataforma Electoral')
    heading(paragraphs[start], spec['seccion'])
    introduction = next((p for p in chapter if p.text.strip() == '[INSERTAR INTRODUCCIÓN]'), None)
    if introduction is None:
        raise ValueError('Falta la introducción de Seguridad.')
    replace_whole(introduction, '[INSERTAR INTRODUCCIÓN]',
                  spec['variables']['seguridad_introduccion_electoral']['marcador'])
    replace_whole(paragraphs[1], '[MUNICIPIO], [ESTADO]', '{municipio}, {estado}')

    for child, title, prefix in zip(children, spec['subsecciones'], ('proteccion_civil', 'seguridad')):
        index = chapter.index(child)
        next_child = next((chapter.index(p) for p in children if chapter.index(p) > index), len(chapter))
        fields = [p for p in chapter[index + 1:next_child] if 'XXXX' in p.text]
        if len(fields) != len(LABELS):
            raise ValueError(f'La subsección {child.text} debe tener cinco campos.')
        for paragraph, (label, suffix) in zip(fields, LABELS):
            replace_field(paragraph, label, spec['variables'][prefix + '_' + suffix]['marcador'])
        heading(child, title)

    keep = {paragraphs[i]._p for i in (0, 1)} | {p._p for p in chapter}
    body = document._element.body
    for node in list(body):
        if node.tag.endswith('}sectPr'):
            continue
        if node not in keep:
            body.remove(node)
    section = next((node for node in body if node.tag.endswith('}sectPr')), None)
    position = list(body).index(section) if section is not None else len(body)
    title_node = deepcopy(children[0]._p)
    body.insert(position, title_node)
    heading(Paragraph(title_node, document), 'Fuentes clave')
    bibliography_node = deepcopy(introduction._p)
    body.insert(position + 1, bibliography_node)
    bibliography = Paragraph(bibliography_node, document)
    bibliography.clear()
    bibliography.add_run('{fuentes_clave}')
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip():
            body.remove(paragraph._p)

    cover_title, identity = document.paragraphs[:2]
    configurar(cover_title._p, 'medicion', 'titulo')
    for run in cover_title.runs:
        run.bold = False
    configurar(identity._p, 'medicion', 'cuerpo')
    for run in identity.runs:
        run.bold = False
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section = document.sections[0]
    usable = section.page_height - section.top_margin - section.bottom_margin
    cover_title.paragraph_format.space_before = Pt(max(0, usable / 12700 / 2 - 55))
    section.different_first_page_header_footer = True
    footer = section.first_page_footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_study = Document(ROOT / 'templates/seguridad_medicion.docx')
    images = [part.blob for part in source_study.part.related_parts.values()
              if getattr(part, 'content_type', '').startswith('image/')]
    if len(images) != 1:
        raise ValueError('No se identificó un logo único de InstitutionWorks en la base de medición.')
    footer.add_run().add_picture(BytesIO(images[0]), width=Cm(5))
    body_paragraphs = document.paragraphs[2:]
    for index, paragraph in enumerate(body_paragraphs):
        value = paragraph.text.strip()
        if value == spec['seccion']:
            role, initial = 'capitulo', True
            paragraph.paragraph_format.page_break_before = True
        elif value in (*spec['subsecciones'], 'Fuentes clave'):
            role, initial = 'subcapitulo', True
        elif value == '{fuentes_clave}':
            role, initial = 'bibliografia', True
        else:
            role = 'cuerpo'
            previous = body_paragraphs[index - 1].text.strip() if index else ''
            initial = previous == spec['seccion'] or previous in spec['subsecciones']
        configurar(paragraph._p, 'medicion', role, inicial=initial)
        if role == 'bibliografia':
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.bold = False
    destination = ROOT / spec['plantilla']
    document.save(destination)
    with zipfile.ZipFile(destination) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    incrustar(files)
    save(destination, files)
    print(destination)


if __name__ == '__main__':
    main()
