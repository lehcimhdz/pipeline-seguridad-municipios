"""Compone y renderiza el apartado de Seguridad de la plataforma electoral."""

import hashlib
import json
from pathlib import Path
import re
import unicodedata
import zipfile

from docx import Document

from redaccion_consultoria import comprobar_texto


ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT / 'config/variables_plataforma_electoral_seguridad.json'
MARKER = re.compile(r'\{([a-z][a-z0-9_]*)\}')


def slug(value):
    value = ''.join(char for char in unicodedata.normalize('NFD', value.lower())
                    if unicodedata.category(char) != 'Mn')
    return re.sub(r'[^a-z0-9]+', '_', value).strip('_')


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def cargar(result, path):
    """Valida un texto revisado y vinculado a las fuentes de la misma ejecución."""
    spec = json.loads(SPEC.read_text(encoding='utf-8'))
    catalog_path = ROOT / spec['referencias']['catalogo']
    catalog = json.loads(catalog_path.read_text(encoding='utf-8'))['fuentes']
    path = Path(path)
    if not path.is_file():
        raise ValueError(f'Falta la redacción electoral del municipio: {path}')
    artifact = json.loads(path.read_text(encoding='utf-8'))
    if artifact.get('municipio') != result['municipio'] or artifact.get('estado') != result['estado']:
        raise ValueError('La identidad de la redacción electoral difiere de la evidencia.')
    if artifact.get('revision', {}).get('estado') != 'revisada':
        raise ValueError('La redacción electoral requiere revisión editorial declarada.')
    if artifact.get('vinculos', {}).get('fuentes') != result['fuentes']:
        raise ValueError('Cambió alguna fuente municipal de la redacción electoral.')
    if artifact['vinculos'].get('evaluacion_sha256') != result['evaluacion_sha256']:
        raise ValueError('Cambió la evaluación municipal: revisar las propuestas electorales.')
    if set(artifact.get('campos', {})) != set(spec['variables']):
        raise ValueError('La redacción electoral no contiene exactamente las variables de contenido previstas.')
    values = {'municipio': result['municipio'], 'estado': result['estado']}
    cited = set()
    for key, entry in artifact['campos'].items():
        if not isinstance(entry, dict):
            raise ValueError(f'Contenido electoral inválido: {key}.')
        value, indicators = entry.get('texto'), entry.get('indicadores')
        references = entry.get('referencias', [])
        if (not isinstance(value, str) or not value.strip()
                or not isinstance(indicators, list) or not indicators
                or any(type(i) is not int or not 1 <= i <= 18 for i in indicators)):
            raise ValueError(f'Falta texto o evidencia por indicador en {key}.')
        if (not isinstance(references, list) or any(not isinstance(ref, str) for ref in references)
                or len(references) != len(set(references))
                or any(ref not in catalog for ref in references)):
            raise ValueError(f'Referencias inválidas o duplicadas en {key}.')
        for ref in references:
            if catalog[ref]['cita'] not in value:
                raise ValueError(f'Falta la cita {catalog[ref]["cita"]} en {key}.')
            if key.startswith('proteccion_civil_') and ref not in ('ley_proteccion_civil', 'marco_sendai'):
                raise ValueError(f'La referencia {ref} no corresponde a Protección Civil.')
            if key.startswith('seguridad_') and key != 'seguridad_introduccion_electoral' and ref in ('ley_proteccion_civil', 'marco_sendai'):
                raise ValueError(f'La referencia {ref} no corresponde a Seguridad Pública.')
            cited.add(ref)
        if any(source['cita'] in value for ref, source in catalog.items() if ref not in references):
            raise ValueError(f'Hay citas no declaradas en {key}.')
        allowed = range(1, 4) if key.startswith('proteccion_civil_') else range(4, 19)
        if key == 'seguridad_introduccion_electoral':
            allowed = range(1, 19)
        if any(i not in allowed for i in indicators):
            raise ValueError(f'El campo {key} cita indicadores de otro apartado.')
        if MARKER.search(value) or '[' in value or ']' in value or 'XXXX' in value:
            raise ValueError(f'El campo {key} conserva instrucciones o marcadores.')
        without_citations = value
        for ref in references:
            without_citations = without_citations.replace(catalog[ref]['cita'], '')
        if re.search(r'(?<!\w)\d+(?:[.,]\d+)?', without_citations):
            raise ValueError(f'El campo {key} contiene cifras que requieren vinculación factual explícita.')
        comprobar_texto(value, key)
        values[key] = value.strip()
    if not {'ley_proteccion_civil', 'marco_sendai', 'ley_seguridad_publica'} <= cited:
        raise ValueError('La plataforma requiere las referencias normativas y el Marco de Sendai en sus argumentos.')
    if not cited.intersection({'braga_focalizacion', 'piza_camaras'}):
        raise ValueError('La justificación de medidas policiales requiere al menos una fuente de investigación.')
    values['fuentes_clave'] = '\n'.join(catalog[ref]['referencia_apa'] for ref in catalog if ref in cited)
    comprobar_texto(values['fuentes_clave'], 'fuentes_clave')
    return {'municipio': result['municipio'], 'estado': result['estado'],
            'evaluacion_sha256': result['evaluacion_sha256'], 'fuentes': result['fuentes'],
            'plantilla_sha256': sha256(ROOT / spec['plantilla']),
            'redaccion_sha256': sha256(path), 'fuentes_clave_sha256': sha256(catalog_path),
            'valores_plantilla': values,
            'referencias': sorted(cited),
            'indicadores_por_variable': {key: artifact['campos'][key]['indicadores']
                                        for key in spec['variables']}}


def renderizar(content, destination):
    spec = json.loads(SPEC.read_text(encoding='utf-8'))
    template = ROOT / spec['plantilla']
    if sha256(template) != content['plantilla_sha256']:
        raise ValueError('La plantilla electoral cambió después de validar su contenido.')
    if 'fuentes_clave_sha256' in content and sha256(ROOT / spec['referencias']['catalogo']) != content['fuentes_clave_sha256']:
        raise ValueError('Cambió el catálogo de fuentes electorales después de validar el contenido.')
    document = Document(template)
    seen = []
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            for marker in MARKER.findall(run.text):
                if marker not in content['valores_plantilla']:
                    raise ValueError(f'Variable electoral desconocida: {marker}.')
                run.text = run.text.replace('{' + marker + '}', content['valores_plantilla'][marker])
                seen.append(marker)
    if set(seen) != set(content['valores_plantilla']) or len(seen) != len(content['valores_plantilla']):
        raise ValueError('Variables electorales faltantes o duplicadas en la plantilla.')
    if any(MARKER.search(paragraph.text) for paragraph in document.paragraphs):
        raise ValueError('La plataforma electoral conserva variables sin llenar.')
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    with zipfile.ZipFile(destination) as archive:
        if archive.testzip() is not None:
            raise ValueError('El Word electoral no superó la revisión de integridad.')
    return {'archivo': str(destination), 'sha256': sha256(destination),
            'variables': len(seen), 'plantilla_sha256': content['plantilla_sha256'],
            'redaccion_sha256': content['redaccion_sha256'],
            'evaluacion_sha256': content['evaluacion_sha256'],
            'fuentes_clave_sha256': content.get('fuentes_clave_sha256')}
